"""Figure provenance and reproducible chemistry DOCX delivery."""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import datetime, timezone
from io import BytesIO
import hashlib
from pathlib import Path
import re
from typing import Iterable
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

try:
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Inches, Pt
except ImportError:  # pragma: no cover - exercised in keyless/minimal installs.
    Document = None  # type: ignore[assignment]
    WD_LINE_SPACING = None  # type: ignore[assignment]
    Inches = Pt = None  # type: ignore[assignment]

try:
    from PIL import Image
except ImportError:  # pragma: no cover - exercised in keyless/minimal installs.
    Image = None  # type: ignore[assignment]

from orchestrator import _document, _section_value, _split_frontmatter
from review import _parse_content_blocks


FIGURE_STATUSES = {"SOURCE", "ADAPTED", "REDRAWN", "GENERATED"}
JOURNAL_PROFILE_STATUSES = {"NOT_SELECTED", "SELECTED"}
JOURNAL_ADAPTATION_STATUSES = {"NOT_APPLICABLE", "MET", "GAP"}
ASSET_TYPES = {"FIGURE", "SCHEME", "TABLE"}
EXTRACTION_STATUSES = {"UNKNOWN", "VERIFIED", "UNVERIFIED", "FAILED"}


class DocxExportError(RuntimeError):
    """Raised when an existing DOCX cannot be safely regenerated."""


@dataclass(frozen=True)
class FigureAsset:
    """One source-bound visual/table asset and its intended placement.

    ``FigureAsset`` is retained as the public name for compatibility with the
    first delivery slice, but it now covers Figures, Schemes, and Tables. A
    non-source asset may be inventoried for an honest work-in-progress record;
    only ``SOURCE`` assets with complete provenance can enter DOCX delivery.
    """

    asset_id: str
    source_id: str
    source_path: str
    locator: str
    caption: str
    provenance: str
    asset_type: str = "FIGURE"
    status: str = "SOURCE"
    sha256: str = ""
    resolution: str = ""
    page: str = ""
    bbox: tuple[int, int, int, int] | str = ""
    extraction_status: str = "UNKNOWN"
    target_section: str = ""
    target_paragraph: str = ""
    claim_ids: tuple[str, ...] = ()
    citation_ids: tuple[str, ...] = ()
    table_rows: tuple[tuple[str, ...], ...] = ()
    transform_history: str = "None"


@dataclass(frozen=True)
class JournalProfile:
    """Versioned journal-format input derived from an official guide snapshot."""

    status: str
    target_journal: str = ""
    guide_locator: str = ""
    guide_retrieved_at: str = ""
    guide_digest: str = ""
    requirements: tuple[str, ...] = ()
    adaptation_status: str = "NOT_APPLICABLE"
    risk: str = ""
    recovery_action: str = ""

    @classmethod
    def unselected(cls) -> "JournalProfile":
        return cls(status="NOT_SELECTED")

    @classmethod
    def selected(
        cls,
        *,
        target_journal: str,
        guide_locator: str,
        guide_retrieved_at: str,
        guide_digest: str,
        requirements: Iterable[str],
        adaptation_status: str = "MET",
        risk: str = "",
        recovery_action: str = "",
    ) -> "JournalProfile":
        values = cls(
            status="SELECTED",
            target_journal=target_journal.strip(),
            guide_locator=guide_locator.strip(),
            guide_retrieved_at=guide_retrieved_at.strip(),
            guide_digest=guide_digest.strip(),
            requirements=tuple(item.strip() for item in requirements if item.strip()),
            adaptation_status=adaptation_status.strip(),
            risk=risk.strip(),
            recovery_action=recovery_action.strip(),
        )
        values.validate()
        return values

    @classmethod
    def from_guide_snapshot(
        cls,
        project_root: str | Path,
        *,
        requirements: Iterable[str] = (),
    ) -> "JournalProfile":
        path = Path(project_root) / "journal-guide.md"
        if not path.exists():
            raise FileNotFoundError("A journal guide snapshot is required for a selected profile.")
        metadata, _ = _split_frontmatter(path.read_text(encoding="utf-8"))
        if metadata.get("kind") != "journal-guide-snapshot":
            raise ValueError("journal-guide.md has the wrong asset kind.")
        content = _section_value(path.read_text(encoding="utf-8"), "Guide content")
        digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
        if digest != metadata.get("content_digest", ""):
            raise ValueError("journal-guide.md content digest does not match its guide content.")
        return cls.selected(
            target_journal=metadata.get("target_journal", ""),
            guide_locator=metadata.get("source_locator", ""),
            guide_retrieved_at=metadata.get("retrieved_at", ""),
            guide_digest=digest,
            requirements=requirements,
            adaptation_status="MET",
        )

    @classmethod
    def reconcile(cls, project_root: str | Path) -> "JournalProfile":
        """Compare a saved profile with the current guide without overwriting it.

        A missing, malformed, or changed guide returns the old profile with a
        ``GAP`` status and an explicit recovery action. This is intentionally a
        pure read/reconciliation operation; callers decide when to persist a
        refreshed profile after human review.
        """

        root = Path(project_root)
        profile_path = root / "journal-profile.md"
        if not profile_path.exists():
            return cls.unselected()
        old = cls._load_persisted(profile_path)
        if old.status == "NOT_SELECTED":
            return old
        guide_path = root / "journal-guide.md"
        if not guide_path.exists():
            return replace(
                old,
                adaptation_status="GAP",
                risk="The official journal guide is unavailable.",
                recovery_action=(
                    "HUMAN_ACTION_REQUIRED: restore or refetch the official guide, "
                    "then review and persist a refreshed profile."
                ),
            )
        try:
            current = cls.from_guide_snapshot(root, requirements=old.requirements)
        except (OSError, ValueError):
            return replace(
                old,
                adaptation_status="GAP",
                risk="The official journal guide is unavailable or invalid.",
                recovery_action=(
                    "HUMAN_ACTION_REQUIRED: restore or refetch the official guide, "
                    "then review and persist a refreshed profile."
                ),
            )
        if (
            current.target_journal != old.target_journal
            or current.guide_locator != old.guide_locator
            or current.guide_digest != old.guide_digest
        ):
            return replace(
                old,
                adaptation_status="GAP",
                risk="The official journal guide changed since the saved profile was captured.",
                recovery_action=(
                    "HUMAN_ACTION_REQUIRED: review the new guide and explicitly persist "
                    "a refreshed journal profile."
                ),
            )
        return replace(old, adaptation_status="MET", risk="", recovery_action="")

    @classmethod
    def _load_persisted(cls, path: Path) -> "JournalProfile":
        text = path.read_text(encoding="utf-8")
        metadata, body = _split_frontmatter(text)
        if metadata.get("kind") != "journal-profile":
            raise ValueError("journal-profile.md has the wrong asset kind.")
        status = metadata.get("status", "NOT_SELECTED")
        if status == "NOT_SELECTED":
            return cls.unselected()
        requirements = []
        for line in _section_value(text, "Requirements").splitlines():
            if line.startswith("- "):
                requirements.append(line[2:].strip())
        return cls.selected(
            target_journal=_section_value(text, "Target journal"),
            guide_locator=_section_value(text, "Official guide"),
            guide_retrieved_at=_section_value(text, "Retrieved at"),
            guide_digest=_section_value(text, "Guide digest") or metadata.get("guide_digest", ""),
            requirements=requirements,
            adaptation_status=metadata.get("adaptation_status", "MET"),
            risk=_section_value(text, "Risk"),
            recovery_action=_section_value(text, "Recovery action"),
        )

    def persist(self, project_root: str | Path) -> Path:
        self.validate()
        path = Path(project_root) / "journal-profile.md"
        body = "# Journal Profile\n\n"
        if self.status == "NOT_SELECTED":
            body += "No target journal selected; generic chemistry formatting remains active.\n"
        else:
            body += (
                f"## Target journal\n{self.target_journal}\n\n"
                f"## Official guide\n{self.guide_locator}\n\n"
                f"## Retrieved at\n{self.guide_retrieved_at}\n\n"
                f"## Guide digest\n{self.guide_digest}\n\n"
                f"## Adaptation status\n{self.adaptation_status}\n\n"
                f"## Risk\n{self.risk or 'None recorded.'}\n\n"
                f"## Recovery action\n{self.recovery_action or 'None recorded.'}\n\n"
                "## Requirements\n"
                + ("\n".join(f"- {item}" for item in self.requirements) or "None recorded.")
                + "\n"
            )
        path.write_text(
            _document(
                {
                    "kind": "journal-profile",
                    "schema": "1",
                    "status": self.status,
                    "target_journal": self.target_journal or "NOT_SELECTED",
                    "guide_digest": self.guide_digest,
                    "adaptation_status": self.adaptation_status,
                },
                body,
            ).rstrip()
            + "\n",
            encoding="utf-8",
        )
        return path

    def validate(self) -> None:
        if self.status not in JOURNAL_PROFILE_STATUSES:
            raise ValueError(f"Unknown journal-profile status: {self.status}")
        if self.status == "NOT_SELECTED":
            if self.adaptation_status != "NOT_APPLICABLE":
                raise ValueError("An unselected journal profile is NOT_APPLICABLE.")
            if any(
                (self.target_journal, self.guide_locator, self.guide_retrieved_at, self.guide_digest)
            ):
                raise ValueError("An unselected journal profile cannot guess a target or guide.")
            return
        if self.adaptation_status not in JOURNAL_ADAPTATION_STATUSES - {"NOT_APPLICABLE"}:
            raise ValueError("A selected journal profile must report MET or GAP.")
        if not all(
            (
                self.target_journal,
                self.guide_locator,
                self.guide_retrieved_at,
                self.guide_digest,
            )
        ):
            raise ValueError("A selected journal profile requires target and official-guide provenance.")


@dataclass(frozen=True)
class DocxExportResult:
    output_path: Path
    manifest_path: Path
    source_digest: str
    output_digest: str
    status: str


class FigureInventory:
    """Project-local Markdown inventory for source-bound visual assets."""

    def __init__(self, project_root: str | Path) -> None:
        self.project_root = Path(project_root)
        self.assets: dict[str, FigureAsset] = {}

    @property
    def path(self) -> Path:
        return self.project_root / "figure-inventory.md"

    def register_source_figure(self, asset: FigureAsset) -> FigureAsset:
        if asset.asset_type != "FIGURE":
            raise ValueError("register_source_figure only accepts FIGURE assets.")
        if asset.status != "SOURCE":
            raise ValueError("Only a verified source figure can enter the source inventory.")
        return self.register_source_asset(asset)

    def register_source_asset(self, asset: FigureAsset) -> FigureAsset:
        """Register a verified source Figure, Scheme, or Table asset."""

        if asset.asset_type not in ASSET_TYPES:
            raise ValueError(f"Unknown visual asset type: {asset.asset_type}")
        if asset.status != "SOURCE":
            raise ValueError("Only a verified source asset can enter the source inventory.")
        if asset.asset_id in self.assets:
            raise ValueError(f"Duplicate figure asset ID: {asset.asset_id}")
        if not all(
            value.strip()
            for value in (
                asset.asset_id,
                asset.source_id,
                asset.locator,
                asset.caption,
                asset.provenance,
            )
        ):
            raise ValueError(
                "A source asset requires identity, locator, caption, and provenance."
            )
        if asset.extraction_status not in EXTRACTION_STATUSES:
            raise ValueError(f"Unknown extraction status: {asset.extraction_status}")
        if _forbidden_source_provenance(asset.provenance):
            raise ValueError(
                "A SOURCE asset cannot claim AI-generated, composite, or redrawn provenance."
            )

        if asset.asset_type == "TABLE":
            if not asset.table_rows:
                raise ValueError("A source table requires at least one table row.")
            table_digest = _table_digest(asset.table_rows)
            if asset.sha256 and asset.sha256 != table_digest:
                raise ValueError(f"Table {asset.asset_id} source hash does not match its rows.")
            registered = replace(
                asset,
                sha256=asset.sha256 or table_digest,
                resolution=asset.resolution or "N/A",
                extraction_status=asset.extraction_status
                if asset.extraction_status != "UNKNOWN"
                else "VERIFIED",
            )
        else:
            if not asset.source_path.strip():
                raise ValueError("A source image asset requires a source path.")
            source_path = self._resolve_source(asset.source_path)
            if not source_path.is_file():
                raise FileNotFoundError(f"Source asset does not exist: {asset.source_path}")
            if Image is None:
                raise RuntimeError("Figure inventory requires Pillow for image metadata.")
            with Image.open(source_path) as image:
                resolution = f"{image.width}x{image.height}"
            source_digest = hashlib.sha256(source_path.read_bytes()).hexdigest()
            if asset.sha256 and asset.sha256 != source_digest:
                raise ValueError(f"Figure {asset.asset_id} source hash does not match its file.")
            registered = replace(
                asset,
                sha256=source_digest,
                resolution=resolution,
                extraction_status=asset.extraction_status
                if asset.extraction_status != "UNKNOWN"
                else "VERIFIED",
            )
        self.assets[asset.asset_id] = registered
        return registered

    def register_asset(self, asset: FigureAsset) -> FigureAsset:
        """Record source or transformed work-in-progress without delivery claims."""

        if asset.asset_type not in ASSET_TYPES:
            raise ValueError(f"Unknown visual asset type: {asset.asset_type}")
        if asset.status == "SOURCE":
            return self.register_source_asset(asset)
        if asset.status not in FIGURE_STATUSES:
            raise ValueError(f"Unknown visual asset status: {asset.status}")
        if asset.asset_id in self.assets:
            raise ValueError(f"Duplicate figure asset ID: {asset.asset_id}")
        if not all(
            value.strip()
            for value in (asset.asset_id, asset.source_id, asset.caption, asset.provenance)
        ):
            raise ValueError("An inventoried asset requires identity, source, caption, and provenance.")
        if asset.extraction_status not in EXTRACTION_STATUSES:
            raise ValueError(f"Unknown extraction status: {asset.extraction_status}")
        self.assets[asset.asset_id] = asset
        return asset

    def place(
        self,
        asset_id: str,
        *,
        section: str,
        paragraph: str = "",
        claim: str = "",
        citation: str = "",
    ) -> FigureAsset:
        if asset_id not in self.assets:
            raise ValueError(f"Unknown figure asset: {asset_id}")
        if not section.strip():
            raise ValueError("Figure placement requires a target section.")
        placed = replace(
            self.assets[asset_id],
            target_section=section.strip(),
            target_paragraph=paragraph.strip(),
            claim_ids=_append_unique(self.assets[asset_id].claim_ids, claim),
            citation_ids=_append_unique(self.assets[asset_id].citation_ids, citation),
        )
        self.assets[asset_id] = placed
        return placed

    def persist(self) -> Path:
        self.project_root.mkdir(parents=True, exist_ok=True)
        body = "# Figure/Scheme/Table Inventory\n\n"
        if not self.assets:
            body += "## Assets\nNone recorded.\n"
        else:
            body += "\n\n".join(_render_figure_asset(asset) for asset in self.assets.values())
        self.path.write_text(
            _document(
                {
                    "kind": "figure-inventory",
                    "schema": "1",
                    "asset_count": str(len(self.assets)),
                },
                body,
            ).rstrip()
            + "\n",
            encoding="utf-8",
        )
        return self.path

    @classmethod
    def load(cls, project_root: str | Path) -> "FigureInventory":
        inventory = cls(project_root)
        if not inventory.path.exists():
            return inventory
        text = inventory.path.read_text(encoding="utf-8")
        metadata, body = _split_frontmatter(text)
        if metadata.get("kind") != "figure-inventory":
            raise ValueError("figure-inventory.md has the wrong asset kind.")
        for match in re.finditer(
            r"^## Asset: (?P<asset_id>.+?)\n(?P<body>.*?)(?=^## Asset: |\Z)",
            body,
            flags=re.MULTILINE | re.DOTALL,
        ):
            values = _parse_labeled_lines(match.group("body"))
            asset = FigureAsset(
                asset_id=match.group("asset_id").strip(),
                source_id=values.get("Source ID", ""),
                source_path=values.get("Source path", ""),
                locator=values.get("Locator", ""),
                caption=values.get("Caption", ""),
                provenance=values.get("Provenance", ""),
                asset_type=values.get("Asset type", "FIGURE"),
                status=values.get("Status", "SOURCE"),
                sha256=values.get("SHA-256", ""),
                resolution=values.get("Resolution", ""),
                page=values.get("Page", ""),
                bbox=values.get("BBox", ""),
                extraction_status=values.get("Extraction status", "UNKNOWN"),
                target_section=values.get("Target section", ""),
                target_paragraph=values.get("Target paragraph", ""),
                claim_ids=_csv_values(values.get("Claim IDs", "")),
                citation_ids=_csv_values(values.get("Citation IDs", "")),
                table_rows=_parse_table_rows(values.get("Table rows", "")),
                transform_history=values.get("Transform history", "None"),
            )
            inventory.assets[asset.asset_id] = asset
        return inventory

    def validate_for_delivery(self) -> tuple[FigureAsset, ...]:
        validated: list[FigureAsset] = []
        for asset in self.assets.values():
            if asset.status != "SOURCE":
                raise ValueError(
                    f"Figure {asset.asset_id} is not a verified source figure: {asset.status}"
                )
            if not all(
                value.strip()
                for value in (
                    asset.source_id,
                    asset.locator,
                    asset.caption,
                    asset.resolution,
                    asset.provenance,
                    asset.target_section,
                    asset.target_paragraph,
                )
            ):
                raise ValueError(
                    f"Figure {asset.asset_id} is missing provenance, locator, or placement."
                )
            if not asset.claim_ids or not asset.citation_ids:
                raise ValueError(
                    f"Figure {asset.asset_id} requires claim and citation placement bindings."
                )
            if asset.asset_type not in ASSET_TYPES:
                raise ValueError(f"Asset {asset.asset_id} has an unknown type: {asset.asset_type}")
            if asset.extraction_status != "VERIFIED":
                raise ValueError(
                    f"Figure {asset.asset_id} is not verified: extraction status {asset.extraction_status}"
                )
            if asset.asset_type == "TABLE":
                if not asset.table_rows:
                    raise ValueError(f"Table {asset.asset_id} has no rows.")
                if asset.sha256 != _table_digest(asset.table_rows):
                    raise ValueError(f"Table {asset.asset_id} source hash has changed.")
            else:
                if not asset.source_path.strip() or not asset.sha256:
                    raise ValueError(f"Figure {asset.asset_id} is missing its source file or hash.")
                source_path = self._resolve_source(asset.source_path)
                if not source_path.is_file():
                    raise FileNotFoundError(f"Source asset does not exist: {asset.source_path}")
                if hashlib.sha256(source_path.read_bytes()).hexdigest() != asset.sha256:
                    raise ValueError(f"Figure {asset.asset_id} source hash has changed.")
            validated.append(asset)
        return tuple(validated)

    def _resolve_source(self, source_path: str) -> Path:
        path = Path(source_path)
        return path if path.is_absolute() else self.project_root / path


class GenericChemistryDocxExporter:
    """Render canonical content and verified figures into a reproducible DOCX."""

    def __init__(self, project_root: str | Path) -> None:
        self.project_root = Path(project_root)

    @property
    def source_path(self) -> Path:
        return self.project_root / "review-content.md"

    def export(
        self,
        output_path: str | Path | None = None,
        *,
        profile: JournalProfile | None = None,
    ) -> DocxExportResult:
        profile = profile or JournalProfile.unselected()
        profile.validate()
        if Document is None:
            raise DocxExportError(
                "DOCX export requires python-docx; use the setup wizard or the Markdown fallback."
            )
        if not self.source_path.exists():
            raise FileNotFoundError("DOCX export requires review-content.md.")
        source_text = self.source_path.read_text(encoding="utf-8")
        source_metadata, _ = _split_frontmatter(source_text)
        if source_metadata.get("kind") != "single-review-content-source":
            raise ValueError("review-content.md has the wrong asset kind.")
        blocks = _parse_content_blocks(source_text)
        source_digest = hashlib.sha256(source_text.encode("utf-8")).hexdigest()
        output = Path(output_path) if output_path is not None else self.project_root / "generic-chemistry-draft.docx"
        if not output.is_absolute():
            output = self.project_root / output
        manifest_path = output.with_suffix(output.suffix + ".manifest.md")
        if output.exists() or manifest_path.exists():
            self.check_existing(output)

        inventory = FigureInventory.load(self.project_root)
        figures = inventory.validate_for_delivery()
        document = self._build_document(blocks, figures, profile)
        qa = _export_qa(blocks, figures, profile)
        payload = _deterministic_docx(document)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(payload)
        output_digest = hashlib.sha256(payload).hexdigest()
        manifest_path.write_text(
            _document(
                {
                    "kind": "docx-export-manifest",
                    "schema": "1",
                    "profile_status": profile.status,
                    "target_journal": profile.target_journal or "NOT_SELECTED",
                    "source_digest": source_digest,
                    "output_digest": output_digest,
                    "figure_count": str(len(figures)),
                    **{key: str(value) for key, value in qa.items()},
                },
                "# DOCX Export Manifest\n\n"
                "## Source\nCanonical review content.\n\n"
                "## Figure assets\n"
                + ("\n".join(f"- {asset.asset_id}" for asset in figures) or "None recorded.")
                + "\n\n"
                "## Export QA\n"
                f"- Status: {qa['status']}\n"
                f"- Headings: {qa['heading_count']}\n"
                f"- References: {qa['reference_count']}\n"
                f"- Figures: {qa['figure_count']}\n"
                f"- Schemes: {qa['scheme_count']}\n"
                f"- Tables: {qa['table_count']}\n"
                f"- Equations: {qa['equation_count']}\n"
                f"- Image resolution: {qa['image_resolution_status']}\n"
                f"- Layout: {qa['layout_status']}\n\n"
                "This is format evidence, not scientific or journal acceptance.\n",
            ).rstrip()
            + "\n",
            encoding="utf-8",
        )
        return DocxExportResult(
            output_path=output,
            manifest_path=manifest_path,
            source_digest=source_digest,
            output_digest=output_digest,
            status="EXPORTED",
        )

    def check_existing(self, output_path: str | Path) -> DocxExportResult:
        output = Path(output_path)
        if not output.is_absolute():
            output = self.project_root / output
        manifest_path = output.with_suffix(output.suffix + ".manifest.md")
        if not output.exists() or not manifest_path.exists():
            raise DocxExportError("DOCX and its export manifest must exist together.")
        metadata, _ = _split_frontmatter(manifest_path.read_text(encoding="utf-8"))
        if metadata.get("kind") != "docx-export-manifest":
            raise DocxExportError("DOCX export manifest has the wrong asset kind.")
        current_source_digest = hashlib.sha256(
            self.source_path.read_text(encoding="utf-8").encode("utf-8")
        ).hexdigest()
        current_output_digest = hashlib.sha256(output.read_bytes()).hexdigest()
        if metadata.get("source_digest") != current_source_digest:
            raise DocxExportError(
                "DOCX conflicts with the current canonical Markdown digest; preserve both and request human action."
            )
        if metadata.get("output_digest") != current_output_digest:
            raise DocxExportError(
                "DOCX has human edits relative to its manifest; preserve the edit and request human action."
            )
        return DocxExportResult(
            output_path=output,
            manifest_path=manifest_path,
            source_digest=current_source_digest,
            output_digest=current_output_digest,
            status="IN_SYNC",
        )

    def _build_document(self, blocks, figures, profile: JournalProfile):
        if Document is None or WD_LINE_SPACING is None or Inches is None or Pt is None:
            raise DocxExportError("DOCX export dependencies are unavailable.")
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        normal = document.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(11)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        document.core_properties.title = (
            f"Chemical Review — {profile.target_journal}"
            if profile.status == "SELECTED"
            else "Chemical Review — Generic Chemistry Draft"
        )
        document.core_properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
        document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
        document.add_heading(document.core_properties.title, level=0)

        figures_by_paragraph: dict[tuple[str, str], list[FigureAsset]] = {}
        for asset in figures:
            figures_by_paragraph.setdefault(
                (asset.target_section, asset.target_paragraph), []
            ).append(asset)
        seen_sections: set[str] = set()
        paragraph_counts: dict[str, int] = {}
        evidence_ids: list[str] = []
        for block in blocks:
            if block.section not in seen_sections:
                document.add_heading(block.section, level=1)
                seen_sections.add(block.section)
            document.add_paragraph(block.text)
            evidence_ids.extend(block.evidence_ids)
            paragraph_counts[block.section] = paragraph_counts.get(block.section, 0) + 1
            paragraph_id = f"P-{paragraph_counts[block.section]}"
            pending = figures_by_paragraph.pop((block.section, paragraph_id), [])
            for asset in pending:
                if asset.asset_type == "TABLE":
                    column_count = max(len(row) for row in asset.table_rows)
                    table = document.add_table(rows=0, cols=column_count)
                    table.style = "Table Grid"
                    for row in asset.table_rows:
                        cells = table.add_row().cells
                        for index, cell in enumerate(cells):
                            cell.text = row[index] if index < len(row) else ""
                else:
                    path = Path(asset.source_path)
                    if not path.is_absolute():
                        path = self.project_root / path
                    picture = _cropped_picture(path, asset.bbox)
                    document.add_picture(picture, width=Inches(5.5))
                caption = document.add_paragraph(f"{asset.asset_id}. {asset.caption}")
                if caption.runs:
                    caption.runs[0].italic = True
                document.add_paragraph(f"Source: {asset.source_id}; {asset.locator}")
        if figures_by_paragraph:
            missing = ", ".join(
                sorted(
                    asset.asset_id
                    for assets in figures_by_paragraph.values()
                    for asset in assets
                )
            )
            raise ValueError(
                f"Assets are placed at paragraphs absent from canonical content: {missing}"
            )
        if evidence_ids:
            document.add_heading("References", level=1)
            for identifier in dict.fromkeys(evidence_ids):
                document.add_paragraph(identifier, style="List Number")
        if profile.status == "SELECTED":
            document.add_heading("Journal profile", level=1)
            document.add_paragraph(
                f"Official guide: {profile.guide_locator}; retrieved {profile.guide_retrieved_at}; "
                f"digest {profile.guide_digest}; adaptation status {profile.adaptation_status}."
            )
            if profile.risk:
                document.add_paragraph(f"Risk: {profile.risk}")
            if profile.recovery_action:
                document.add_paragraph(f"Recovery action: {profile.recovery_action}")
            for requirement in profile.requirements:
                document.add_paragraph(requirement, style="List Bullet")
        return document


def _export_qa(blocks, assets: tuple[FigureAsset, ...], profile: JournalProfile) -> dict[str, object]:
    evidence_ids: list[str] = []
    for block in blocks:
        evidence_ids.extend(block.evidence_ids)
    references = tuple(dict.fromkeys(evidence_ids))
    sections = tuple(dict.fromkeys(block.section for block in blocks))
    image_assets = tuple(asset for asset in assets if asset.asset_type != "TABLE")
    return {
        "status": "MET",
        "asset_count": len(assets),
        "heading_count": 1 + len(sections) + bool(references) + (profile.status == "SELECTED"),
        "reference_count": len(references),
        "citation_count": len(references),
        "figure_count": sum(asset.asset_type == "FIGURE" for asset in assets),
        "scheme_count": sum(asset.asset_type == "SCHEME" for asset in assets),
        "table_count": sum(asset.asset_type == "TABLE" for asset in assets),
        "equation_count": sum(
            len(re.findall(r"(?im)^\s*(?:Equation|Eq\.?):", block.text)) for block in blocks
        ),
        "image_resolution_status": "MET"
        if all(asset.resolution.strip() for asset in image_assets)
        else "GAP",
        "layout_status": "MET",
    }


def _render_figure_asset(asset: FigureAsset) -> str:
    table_rows = "; ".join(" || ".join(cell for cell in row) for row in asset.table_rows)
    bbox = _format_bbox(asset.bbox)
    return (
        f"## Asset: {asset.asset_id}\n"
        f"Asset type: {asset.asset_type}\n"
        f"Status: {asset.status}\n"
        f"Source ID: {asset.source_id}\n"
        f"Source path: {asset.source_path}\n"
        f"Locator: {asset.locator}\n"
        f"Page: {asset.page}\n"
        f"BBox: {bbox}\n"
        f"Caption: {asset.caption}\n"
        f"SHA-256: {asset.sha256}\n"
        f"Resolution: {asset.resolution}\n"
        f"Extraction status: {asset.extraction_status}\n"
        f"Target section: {asset.target_section}\n"
        f"Target paragraph: {asset.target_paragraph}\n"
        f"Claim IDs: {', '.join(asset.claim_ids)}\n"
        f"Citation IDs: {', '.join(asset.citation_ids)}\n"
        f"Table rows: {table_rows}\n"
        f"Transform history: {asset.transform_history}\n"
        f"Provenance: {asset.provenance}\n"
    )


def _parse_labeled_lines(value: str) -> dict[str, str]:
    parsed: dict[str, str] = {}
    for line in value.splitlines():
        label, separator, item = line.partition(":")
        if separator:
            parsed[label.strip()] = item.strip()
    return parsed


def _csv_values(value: str) -> tuple[str, ...]:
    return tuple(item.strip() for item in value.split(",") if item.strip())


def _parse_table_rows(value: str) -> tuple[tuple[str, ...], ...]:
    if not value.strip():
        return ()
    rows: list[tuple[str, ...]] = []
    for raw_row in value.split(";"):
        cells = tuple(cell.strip() for cell in raw_row.split("||"))
        if any(cells):
            rows.append(cells)
    return tuple(rows)


def _format_bbox(value: tuple[int, int, int, int] | str) -> str:
    if isinstance(value, tuple):
        return ",".join(str(item) for item in value)
    return str(value).strip()


def _append_unique(values: tuple[str, ...], value: str) -> tuple[str, ...]:
    value = value.strip()
    if not value or value in values:
        return values
    return values + (value,)


def _table_digest(rows: tuple[tuple[str, ...], ...]) -> str:
    canonical = "\n".join("\t".join(cell.strip() for cell in row) for row in rows)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _forbidden_source_provenance(value: str) -> bool:
    normalized = re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()
    return bool(
        re.search(
            r"(?:\bai generated\b|\bgenerated by\b|\bsynthetic\b|\bcomposite\b|\bredrawn?\b)",
            normalized,
        )
    )


def _cropped_picture(path: Path, bbox: tuple[int, int, int, int] | str):
    values = _bbox_values(bbox)
    if values is None:
        return str(path)
    if Image is None:
        raise DocxExportError("Figure cropping requires Pillow.")
    with Image.open(path) as source:
        left, top, right, bottom = values
        if not (0 <= left < right <= source.width and 0 <= top < bottom <= source.height):
            raise ValueError(f"Figure crop is outside the source image: {path.name}")
        cropped = source.crop(values)
        stream = BytesIO()
        cropped.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _bbox_values(value: tuple[int, int, int, int] | str) -> tuple[int, int, int, int] | None:
    if isinstance(value, tuple):
        return value
    if not value.strip():
        return None
    parts = tuple(part.strip() for part in value.split(","))
    if len(parts) != 4 or not all(part.isdigit() for part in parts):
        raise ValueError("Figure BBox must contain four non-negative integers.")
    return tuple(int(part) for part in parts)  # type: ignore[return-value]


def _deterministic_docx(document) -> bytes:
    raw = BytesIO()
    document.save(raw)
    source = BytesIO(raw.getvalue())
    target = BytesIO()
    with ZipFile(source, "r") as archive, ZipFile(target, "w", compression=ZIP_DEFLATED) as output:
        for name in sorted(archive.namelist()):
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            output.writestr(info, archive.read(name))
    return target.getvalue()
