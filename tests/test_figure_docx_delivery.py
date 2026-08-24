from __future__ import annotations

import hashlib
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest

from PIL import Image
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SKILL_DIR = ROOT / ".agents" / "skills" / "chemical-review"
sys.path.insert(0, str(SKILL_DIR))

from delivery import (  # noqa: E402
    DocxExportError,
    GenericChemistryDocxExporter,
    FigureAsset,
    FigureInventory,
    JournalProfile,
)
from orchestrator import ChemicalReviewOrchestrator, _document, _split_frontmatter  # noqa: E402


class FigureDocxDeliveryTests(unittest.TestCase):
    @staticmethod
    def _write_source_registry(
        root: Path,
        identity: str,
        *,
        digest: str = "none",
        media_ids: str = "none",
    ) -> None:
        root.joinpath("source-registry.md").write_text(
            "---\nkind: research-source-registry\nschema: 1\n---\n\n"
            "| Source ID | Title | Identity | Kind | Provider | Local path | Access basis | Priority | Metadata | Full text | Parser | Locator(s) | Media IDs | Digest | Failure/recovery |\n"
            "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |\n"
            f"| paper:fixture | Fixture | {identity} | DISCOVERED_METADATA | fixture | none | OPEN_ACCESS | NORMAL | DISCOVERED | FOUND | PARSED | p. 1 | {media_ids} | {digest} | none |\n",
            encoding="utf-8",
        )

    def test_inventory_supports_scheme_table_and_bound_placement(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            scheme_path = root / "scheme.png"
            Image.new("RGB", (96, 64), "white").save(scheme_path)
            inventory = FigureInventory(root)
            scheme = inventory.register_source_asset(
                FigureAsset(
                    asset_id="scheme-1",
                    asset_type="SCHEME",
                    source_id="paper-1",
                    source_path="scheme.png",
                    locator="p. 6, Scheme 1",
                    page="6",
                    bbox=(1, 2, 90, 50),
                    caption="Reaction scheme.",
                    provenance="Cropped from the cited paper.",
                    claim_ids=("claim-1",),
                    citation_ids=("paper-1",),
                    extraction_status="VERIFIED",
                )
            )
            table = inventory.register_source_asset(
                FigureAsset(
                    asset_id="table-1",
                    asset_type="TABLE",
                    source_id="paper-1",
                    source_path="",
                    locator="p. 7, Table 2",
                    caption="Reported conditions.",
                    provenance="Transcribed from the cited paper.",
                    table_rows=(("Entry", "Yield"), ("A", "84%")),
                )
            )
            inventory.place(
                scheme.asset_id,
                section="Mechanistic comparison",
                paragraph="P-2",
                claim="claim-1",
                citation="paper-1",
            )
            inventory.place(
                table.asset_id,
                section="Mechanistic comparison",
                paragraph="P-1",
                claim="claim-1",
                citation="paper-1",
            )
            path = inventory.persist()

            loaded = FigureInventory.load(root)
            self.assertEqual(loaded.assets["scheme-1"].asset_type, "SCHEME")
            self.assertEqual(loaded.assets["scheme-1"].bbox, "1,2,90,50")
            self.assertEqual(loaded.assets["scheme-1"].claim_ids, ("claim-1",))
            self.assertEqual(loaded.assets["table-1"].table_rows[1], ("A", "84%"))
            text = path.read_text(encoding="utf-8")
            self.assertIn("Extraction status: VERIFIED", text)
            self.assertIn("Target paragraph: P-2", text)
            self.assertIn("Claim IDs: claim-1", text)
            self.assertIn("Citation IDs: paper-1", text)

    def test_non_source_asset_can_be_inventoried_but_delivery_rejects_it(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            inventory = FigureInventory(root)
            inventory.register_asset(
                FigureAsset(
                    asset_id="generated-1",
                    source_id="model",
                    source_path="generated.png",
                    locator="",
                    caption="Generated.",
                    provenance="AI generated; not a source figure.",
                    status="GENERATED",
                    extraction_status="UNVERIFIED",
                )
            )
            inventory.persist()
            with self.assertRaisesRegex(ValueError, "not a verified source"):
                FigureInventory.load(root).validate_for_delivery()

    def test_generic_docx_renders_table_and_manifest_qa(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            inventory = FigureInventory(root)
            inventory.register_source_asset(
                FigureAsset(
                    asset_id="table-1",
                    asset_type="TABLE",
                    source_id="paper-1",
                    source_path="",
                    locator="p. 7, Table 2",
                    caption="Reported conditions.",
                    provenance="Transcribed from the cited paper.",
                    table_rows=(("Entry", "Yield"), ("A", "84%")),
                    target_section="Results",
                    target_paragraph="P-1",
                    claim_ids=("Merge 1 · Block",),
                    citation_ids=("paper-1",),
                    extraction_status="VERIFIED",
                )
            )
            inventory.persist()
            self._write_source_registry(root, "paper-1")
            root.joinpath("review-content.md").write_text(
                _document(
                    {
                        "kind": "single-review-content-source",
                        "schema": "1",
                        "content_revision": "1",
                    },
                    "# Review Content Source\n\n"
                    "## Content blocks\n\n"
                    "### Merge 1 · Block\n"
                    "Section: Results\n"
                    "Claim level: SOURCE_FACT\n"
                    "Contribution type: comparison\n"
                    "Source units: unit-1\n"
                    "Evidence IDs: paper-1\n"
                    "Text:\nThe reported yield is 84%.\n",
                ),
                encoding="utf-8",
            )

            exported = GenericChemistryDocxExporter(root).export()
            document = Document(exported.output_path)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(1, 1).text, "84%")
            manifest = exported.manifest_path.read_text(encoding="utf-8")
            self.assertIn("table_count: 1", manifest)
            self.assertIn("reference_count: 1", manifest)
            self.assertIn("layout_status: MANUAL_REVIEW_REQUIRED", manifest)
            self.assertIn("## Export QA", manifest)

    def test_docx_rejects_figure_claim_id_from_a_different_content_block(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            inventory = FigureInventory(root)
            inventory.register_source_asset(
                FigureAsset(
                    asset_id="table-1",
                    asset_type="TABLE",
                    source_id="paper-1",
                    source_path="",
                    locator="p. 7, Table 2",
                    caption="Reported conditions.",
                    provenance="Transcribed from the cited paper.",
                    table_rows=(("Entry", "Yield"), ("A", "84%")),
                    target_section="Results",
                    target_paragraph="P-1",
                    claim_ids=("Merge 1 · Block 2",),
                    citation_ids=("paper-1",),
                    extraction_status="VERIFIED",
                )
            )
            inventory.persist()
            self._write_source_registry(root, "paper-1")
            self._write_two_block_content(root)

            with self.assertRaisesRegex(ValueError, "claim ID.*target paragraph"):
                GenericChemistryDocxExporter(root).export()

    def test_docx_rejects_citation_found_elsewhere_but_not_in_target_paragraph(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            inventory = FigureInventory(root)
            inventory.register_source_asset(
                FigureAsset(
                    asset_id="table-1",
                    asset_type="TABLE",
                    source_id="paper-1",
                    source_path="",
                    locator="p. 7, Table 2",
                    caption="Reported conditions.",
                    provenance="Transcribed from the cited paper.",
                    table_rows=(("Entry", "Yield"), ("A", "84%")),
                    target_section="Results",
                    target_paragraph="P-2",
                    claim_ids=("Merge 1 · Block 2",),
                    citation_ids=("paper-1",),
                    extraction_status="VERIFIED",
                )
            )
            inventory.persist()
            self._write_source_registry(root, "paper-1")
            self._write_two_block_content(root)

            with self.assertRaisesRegex(ValueError, "citation.*target paragraph"):
                GenericChemistryDocxExporter(root).export()

    def test_orchestrator_rejects_direct_export_without_canonical_workflow_state(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            root.joinpath("review-content.md").write_text(
                _document(
                    {
                        "kind": "single-review-content-source",
                        "schema": "1",
                        "content_revision": "1",
                    },
                    "# Review Content Source\n\n"
                    "## Content blocks\n\n"
                    "### Merge 1 · Block\n"
                    "Section: Background\n"
                    "Claim level: MODEL_SYNTHESIS\n"
                    "Contribution type: explanation\n"
                    "Source units: unit-1\n"
                    "Evidence IDs: paper-1\n"
                    "Text:\nA bounded synthesis.\n",
                ),
                encoding="utf-8",
            )
            orchestrator = ChemicalReviewOrchestrator(root)
            with self.assertRaisesRegex(DocxExportError, "canonical delivery-ready"):
                orchestrator.export_docx()

    def test_source_figure_inventory_preserves_locator_hash_and_placement(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            source_image = root / "source-figure.png"
            Image.new("RGB", (120, 80), "white").save(source_image)

            inventory = FigureInventory(root)
            asset = inventory.register_source_figure(
                FigureAsset(
                    asset_id="fig-1",
                    source_id="https://doi.org/10.1000/example",
                    source_path="source-figure.png",
                    locator="p. 4, Figure 2",
                    caption="Source caption.",
                    provenance="Original figure from the cited paper.",
                )
            )
            inventory.place(
                "fig-1",
                section="Mechanistic comparison",
                paragraph="P-2",
                claim="claim-1",
                citation="https://doi.org/10.1000/example",
            )
            path = inventory.persist()

            text = path.read_text(encoding="utf-8")
            metadata, _ = _split_frontmatter(text)
            self.assertEqual(metadata["kind"], "figure-inventory")
            self.assertIn("fig-1", text)
            self.assertIn("p. 4, Figure 2", text)
            self.assertIn("Mechanistic comparison", text)
            self.assertIn(hashlib.sha256(source_image.read_bytes()).hexdigest(), text)
            self.assertTrue(asset.resolution)

    def test_source_pdf_digest_is_distinct_from_asset_hash_and_registry_bound(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            source_image = root / "source-figure.png"
            Image.new("RGB", (120, 80), "white").save(source_image)
            inventory = FigureInventory(root)
            registered = inventory.register_source_figure(
                FigureAsset(
                    asset_id="fig-1",
                    source_id="paper-1",
                    source_path=source_image.name,
                    locator="p. 4, Figure 2",
                    caption="Source caption.",
                    provenance="Original figure from the cited paper.",
                    source_digest="source-pdf-digest",
                    extraction_status="VERIFIED",
                    target_section="Results",
                    target_paragraph="P-1",
                    claim_ids=("claim-1",),
                    citation_ids=("paper-1",),
                )
            )
            inventory.persist()

            loaded = FigureInventory.load(root).assets["fig-1"]
            self.assertEqual(loaded.source_digest, "source-pdf-digest")
            self.assertEqual(
                registered.sha256,
                hashlib.sha256(source_image.read_bytes()).hexdigest(),
            )
            self.assertNotEqual(registered.sha256, registered.source_digest)

            self._write_source_registry(
                root,
                "paper-1",
                digest="different-source-digest",
                media_ids="fig-1",
            )
            with self.assertRaisesRegex(ValueError, "source digest"):
                FigureInventory.load(root).validate_for_delivery()

    def test_non_source_figure_cannot_enter_source_inventory(self):
        with TemporaryDirectory() as project_dir:
            with self.assertRaisesRegex(ValueError, "source figure"):
                FigureInventory(Path(project_dir)).register_source_figure(
                    FigureAsset(
                        asset_id="generated-1",
                        source_id="model",
                        source_path="generated.png",
                        locator="",
                        caption="Generated.",
                        provenance="AI generated",
                        status="GENERATED",
                    )
                )

    def test_source_registration_rejects_transformed_provenance_and_hash_mismatch(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            image_path = root / "source.png"
            Image.new("RGB", (32, 24), "white").save(image_path)
            inventory = FigureInventory(root)
            with self.assertRaisesRegex(ValueError, "AI-generated"):
                inventory.register_source_figure(
                    FigureAsset(
                        asset_id="ai-pretender",
                        source_id="paper-1",
                        source_path=image_path.name,
                        locator="p. 1, Figure 1",
                        caption="Pretended source.",
                        provenance="AI-generated composite based on the source.",
                    )
                )
            with self.assertRaisesRegex(ValueError, "hash"):
                inventory.register_source_figure(
                    FigureAsset(
                        asset_id="hash-mismatch",
                        source_id="paper-1",
                        source_path=image_path.name,
                        locator="p. 1, Figure 1",
                        caption="Mismatched source.",
                        provenance="Cropped from the cited paper.",
                        sha256="0" * 64,
                    )
                )

    def test_generic_docx_rebuilds_from_canonical_content_with_source_figure(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            image_path = root / "source-figure.png"
            Image.new("RGB", (100, 160), "white").save(image_path)
            inventory = FigureInventory(root)
            inventory.register_source_figure(
                FigureAsset(
                    asset_id="fig-1",
                    source_id="paper-1",
                    source_path="source-figure.png",
                    locator="p. 4, Figure 2",
                    caption="Mechanistic source figure.",
                    provenance="Original source figure.",
                    target_section="Mechanistic comparison",
                    target_paragraph="P-1",
                    claim_ids=("Merge 1 · Block",),
                    citation_ids=("paper-1",),
                    extraction_status="VERIFIED",
                )
            )
            inventory.persist()
            self._write_source_registry(root, "paper-1")
            root.joinpath("review-content.md").write_text(
                _document(
                    {
                        "kind": "single-review-content-source",
                        "schema": "1",
                        "content_revision": "1",
                    },
                    "# Review Content Source\n\n"
                    "## Content blocks\n\n"
                    "### Merge 1 · Block\n"
                    "Section: Mechanistic comparison\n"
                    "Claim level: MODEL_SYNTHESIS\n"
                    "Contribution type: comparison\n"
                    "Source units: unit-1\n"
                    "Evidence IDs: paper-1\n"
                    "Text:\nThe source figure supports the comparison.\n",
                ),
                encoding="utf-8",
            )

            first = GenericChemistryDocxExporter(root).export()
            second_path = root / "second.docx"
            second = GenericChemistryDocxExporter(root).export(second_path)

            self.assertEqual(first.source_digest, second.source_digest)
            self.assertEqual(first.output_path.read_bytes(), second.output_path.read_bytes())
            document = Document(first.output_path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("The source figure supports the comparison.", text)
            self.assertIn("Mechanistic source figure.", text)
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertLessEqual(document.inline_shapes[0].height.inches, 5.0)

    @staticmethod
    def _write_two_block_content(root):
        root.joinpath("review-content.md").write_text(
            _document(
                {
                    "kind": "single-review-content-source",
                    "schema": "1",
                    "content_revision": "1",
                },
                "# Review Content Source\n\n"
                "## Content blocks\n\n"
                "### Merge 1 · Block 1\n"
                "Section: Results\n"
                "Claim level: SOURCE_FACT\n"
                "Contribution type: comparison\n"
                "Source units: unit-1\n"
                "Evidence IDs: paper-1\n"
                "Text:\nPaper one result.\n\n"
                "### Merge 1 · Block 2\n"
                "Section: Results\n"
                "Claim level: SOURCE_FACT\n"
                "Contribution type: comparison\n"
                "Source units: unit-2\n"
                "Evidence IDs: paper-2\n"
                "Text:\nPaper two result.\n",
            ),
            encoding="utf-8",
        )

    def test_docx_digest_conflict_preserves_edit_and_pauses(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            root.joinpath("review-content.md").write_text(
                _document(
                    {
                        "kind": "single-review-content-source",
                        "schema": "1",
                        "content_revision": "1",
                    },
                    "# Review Content Source\n\n"
                    "## Content blocks\n\n"
                    "### Merge 1 · Block\n"
                    "Section: Background\n"
                    "Claim level: SOURCE_FACT\n"
                    "Contribution type: term_verification\n"
                    "Source units: unit-1\n"
                    "Evidence IDs: paper-1\n"
                    "Text:\nA source fact.\n",
                ),
                encoding="utf-8",
            )
            exported = GenericChemistryDocxExporter(root).export()
            result = GenericChemistryDocxExporter(root).check_existing(exported.output_path)
            self.assertEqual(result.status, "IN_SYNC")

            exported.output_path.with_suffix(".docx").write_bytes(
                exported.output_path.read_bytes() + b"manual edit"
            )
            with self.assertRaises(DocxExportError):
                GenericChemistryDocxExporter(root).check_existing(exported.output_path)

    def test_unselected_journal_profile_does_not_guess_target(self):
        profile = JournalProfile.unselected()
        self.assertEqual(profile.status, "NOT_SELECTED")
        self.assertEqual(profile.target_journal, "")

    def test_selected_journal_profile_persists_official_guide_provenance(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            guide_content = "Current official review requirements."
            guide_digest = hashlib.sha256(guide_content.encode("utf-8")).hexdigest()
            guide = root / "journal-guide.md"
            guide.write_text(
                _document(
                    {
                        "kind": "journal-guide-snapshot",
                        "schema": "1",
                        "target_journal": "Example Chemistry",
                        "source_locator": "https://example.org/guide",
                        "retrieved_at": "2026-08-24",
                        "content_digest": guide_digest,
                    },
                    "# Official Journal Guide Snapshot\n\n"
                    f"## Guide content\n{guide_content}\n",
                ),
                encoding="utf-8",
            )

            profile = JournalProfile.from_guide_snapshot(
                root,
                requirements=("Narrative review structure",),
            )
            profile_path = profile.persist(root)
            metadata, body = _split_frontmatter(profile_path.read_text(encoding="utf-8"))

            self.assertEqual(metadata["kind"], "journal-profile")
            self.assertEqual(metadata["status"], "SELECTED")
            self.assertIn("Example Chemistry", body)
            self.assertIn(guide_digest, body)

    def test_journal_guide_format_constraints_drive_reproducible_docx_styles(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            guide_content = (
                "Margins: 2 cm.\n"
                "Font: Arial 10 pt.\n"
                "Single line spacing.\n"
                "Maximum 5,000 words.\n"
                "References must use the journal style."
            )
            guide_digest = hashlib.sha256(guide_content.encode("utf-8")).hexdigest()
            root.joinpath("journal-guide.md").write_text(
                _document(
                    {
                        "kind": "journal-guide-snapshot",
                        "schema": "1",
                        "target_journal": "Example Chemistry",
                        "source_locator": "https://example.org/guide",
                        "retrieved_at": "2026-08-24",
                        "content_digest": guide_digest,
                    },
                    "# Official Journal Guide Snapshot\n\n"
                    f"## Guide content\n{guide_content}",
                ),
                encoding="utf-8",
            )
            profile = JournalProfile.from_guide_snapshot(root)
            self.assertEqual(profile.adaptation_status, "GAP")
            self.assertEqual(
                profile.mapped_requirements,
                ("Margins: 2 cm.", "Font: Arial 10 pt.", "Single line spacing."),
            )
            self.assertEqual(
                profile.unmapped_requirements,
                ("Maximum 5,000 words.", "References must use the journal style."),
            )
            root.joinpath("review-content.md").write_text(
                _document(
                    {
                        "kind": "single-review-content-source",
                        "schema": "1",
                        "content_revision": "1",
                    },
                    "# Review Content Source\n\n"
                    "## Content blocks\n\n"
                    "### Merge 1 · Block\n"
                    "Section: Results\n"
                    "Claim level: MODEL_SYNTHESIS\n"
                    "Contribution type: explanation\n"
                    "Source units: unit-1\n"
                    "Evidence IDs: paper-1\n"
                    "Text:\nA bounded synthesis.\n",
                ),
                encoding="utf-8",
            )

            exported = GenericChemistryDocxExporter(root).export(profile=profile)
            document = Document(exported.output_path)
            section = document.sections[0]
            self.assertAlmostEqual(section.left_margin.inches, 2 / 2.54, places=3)
            self.assertEqual(document.styles["Normal"].font.name, "Arial")
            self.assertAlmostEqual(document.styles["Normal"].font.size.pt, 10, places=2)
            manifest = exported.manifest_path.read_text(encoding="utf-8")
            self.assertIn("journal_format_mapping: PARTIAL_GAP", manifest)

    def test_changed_or_unavailable_guide_preserves_old_profile_and_requests_recovery(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            old_content = "Old official review requirements."
            old_digest = hashlib.sha256(old_content.encode("utf-8")).hexdigest()
            guide = root / "journal-guide.md"
            guide.write_text(
                _document(
                    {
                        "kind": "journal-guide-snapshot",
                        "schema": "1",
                        "target_journal": "Example Chemistry",
                        "source_locator": "https://example.org/guide",
                        "retrieved_at": "2026-08-24",
                        "content_digest": old_digest,
                    },
                    "# Official Journal Guide Snapshot\n\n"
                    f"## Guide content\n{old_content}\n",
                ),
                encoding="utf-8",
            )
            original = JournalProfile.from_guide_snapshot(
                root,
                requirements=("Narrative review structure",),
            )
            original.persist(root)

            changed_content = "Changed official review requirements."
            guide.write_text(
                _document(
                    {
                        "kind": "journal-guide-snapshot",
                        "schema": "1",
                        "target_journal": "Example Chemistry",
                        "source_locator": "https://example.org/guide",
                        "retrieved_at": "2026-08-25",
                        "content_digest": hashlib.sha256(
                            changed_content.encode("utf-8")
                        ).hexdigest(),
                    },
                    "# Official Journal Guide Snapshot\n\n"
                    f"## Guide content\n{changed_content}\n",
                ),
                encoding="utf-8",
            )
            reconciled = JournalProfile.reconcile(root)
            self.assertEqual(reconciled.target_journal, original.target_journal)
            self.assertEqual(reconciled.guide_digest, original.guide_digest)
            self.assertEqual(reconciled.adaptation_status, "GAP")
            self.assertIn("changed", reconciled.risk.lower())
            self.assertIn("HUMAN_ACTION_REQUIRED", reconciled.recovery_action)

            guide.unlink()
            unavailable = JournalProfile.reconcile(root)
            self.assertEqual(unavailable.guide_digest, original.guide_digest)
            self.assertEqual(unavailable.adaptation_status, "GAP")
            self.assertIn("unavailable", unavailable.risk.lower())

    def test_unchanged_guide_cannot_resurrect_an_unmapped_profile_as_met(self):
        with TemporaryDirectory() as project_dir:
            root = Path(project_dir)
            guide_content = "References must use the journal style."
            digest = hashlib.sha256(guide_content.encode("utf-8")).hexdigest()
            root.joinpath("journal-guide.md").write_text(
                _document(
                    {
                        "kind": "journal-guide-snapshot",
                        "schema": "1",
                        "target_journal": "Example Chemistry",
                        "source_locator": "https://example.org/guide",
                        "retrieved_at": "2026-08-24",
                        "content_digest": digest,
                    },
                    "# Official Journal Guide Snapshot\n\n"
                    f"## Guide content\n{guide_content}\n",
                ),
                encoding="utf-8",
            )
            JournalProfile.from_guide_snapshot(root).persist(root)

            reconciled = JournalProfile.reconcile(root)

            self.assertEqual(reconciled.adaptation_status, "GAP")
            self.assertEqual(reconciled.unmapped_requirements, (guide_content,))
            persisted = JournalProfile._load_persisted(root / "journal-profile.md")
            self.assertEqual(persisted.adaptation_status, "GAP")


if __name__ == "__main__":
    unittest.main()
